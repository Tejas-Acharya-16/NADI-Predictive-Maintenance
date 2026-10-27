"""Build the final FD001 RUL implementation report as a DOCX source for PDF export."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
REPORTS = ROOT / "reports"
OUTPUT = REPORTS / "Final_FD001_RUL_Implementation_Report.docx"
FIGURE = REPORTS / "figures" / "ablation_comparison.png"

NAVY = "17365D"
BLUE = "D9EAF7"
PALE = "F4F7FA"
GRID = "D9D9D9"
TEXT = RGBColor(0, 0, 0)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def border_table(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), GRID)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    tr_pr.append(hdr)


def cell_text(cell, text: str, *, bold: bool = False, center: bool = False, color=TEXT) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(8.5)
    run.font.color.rgb = color
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side in ("top", "start", "bottom", "end"):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), "95")
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    tc_pr.append(margins)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    border_table(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        header.cells[index].width = Cm(widths[index])
        shade(header.cells[index], NAVY)
        cell_text(header.cells[index], value, bold=True, center=True, color=RGBColor(255, 255, 255))
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].width = Cm(widths[index])
            if row_index % 2 == 1:
                shade(cells[index], PALE)
            cell_text(cells[index], value, center=index > 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def para(doc: Document, text: str, *, bold_start: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    if bold_start and text.startswith(bold_start):
        run = p.add_run(bold_start)
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(10)
        tail = p.add_run(text[len(bold_start) :])
        tail.font.name = "Aptos"
        tail.font.size = Pt(10)
    else:
        run = p.add_run(text)
        run.font.name = "Aptos"
        run.font.size = Pt(10)


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(9.5)


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    run.font.name = "Aptos Display" if level == 1 else "Aptos"
    run.font.color.rgb = TEXT
    run.font.size = Pt(15 if level == 1 else 11.5)
    run.bold = True


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("FD001 RUL Implementation Report   |   Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.75)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(1.85)
    section.right_margin = Cm(1.85)
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "Aptos"
        style.font.color.rgb = TEXT
    doc.styles["Normal"].font.size = Pt(10)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(89, 89, 89)


def main() -> None:
    doc = Document()
    configure(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("NASA C MAPSS FD001 Remaining Useful Life")
    run.font.name = "Aptos Display"
    run.font.size = Pt(25)
    run.font.bold = True
    run.font.color.rgb = TEXT

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.paragraph_format.space_after = Pt(16)
    run = subtitle.add_run("Final implementation, model comparison, and deployment assessment")
    run.font.name = "Aptos"
    run.font.size = Pt(12.5)
    run.font.color.rgb = RGBColor(89, 89, 89)

    info = doc.add_paragraph()
    info.paragraph_format.space_after = Pt(14)
    run = info.add_run(f"Prepared {date.today().strftime('%d %B %Y')}   |   CPU-bounded local study")
    run.font.name = "Aptos"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(89, 89, 89)

    heading(doc, "Executive summary")
    para(
        doc,
        "The selected production candidate is a Histogram Gradient Boosting regressor trained with RUL labels capped at 125 cycles and causal 10- and 30-cycle features. It was selected on held-out engine prefixes and then retrained on all FD001 training engines. Its single final official-test evaluation achieved RMSE 13.532, MAE 9.799, R² 0.894, and PHM08 mean 2.852.",
    )
    para(
        doc,
        "The study also implements compact CNN, LSTM, CNN-LSTM, TinyML-oriented, autoencoder, hybrid, random forest, ridge, and mean-baseline paths. The compact CNN-LSTM was the best neural candidate, but it did not outperform the selected tree model. The autoencoder is retained as an unlabeled early-life anomaly signal for triage rather than as a supervised RUL predictor.",
    )

    add_table(
        doc,
        ["Stage", "RMSE", "MAE", "R²", "PHM08 mean"],
        [
            ["Validation selection split", "12.773", "8.609", "0.877", "2.832"],
            ["Three-split validation mean ± SD", "13.623 ± 1.724", "8.802 ± 0.561", "0.866 ± 0.025", "3.471 ± 1.411"],
            ["Official FD001 test, one final evaluation", "13.532", "9.799", "0.894", "2.852"],
        ],
        [7.1, 2.3, 2.3, 2.0, 2.8],
    )

    heading(doc, "1  Problem and data")
    para(
        doc,
        "The task predicts remaining useful life from multivariate turbofan telemetry. The FD001 files contain a unit identifier, cycle, three operating settings, and 21 sensor channels. Training trajectories run to failure; test trajectories are right-censored and have one supplied RUL target for each engine endpoint.",
    )
    add_table(
        doc,
        ["Split", "Rows", "Engines", "Observed life / cycles", "Audit result"],
        [
            ["Training", "20,631", "100", "128 to 362 cycles", "All quality checks passed"],
            ["Test", "13,096", "100", "31 to 303 observed cycles", "All quality checks passed"],
            ["Test RUL vector", "100", "100", "7 to 145 cycles", "Aligned to sorted engine IDs"],
        ],
        [3.0, 2.0, 1.9, 4.3, 5.3],
    )
    para(doc, "No missing or non-finite values, duplicate unit-cycle rows, cycle gaps, invalid IDs, or ordering faults were found. Seven near-constant sensors were screened in an ablation, but removing them did not improve the selected result.")

    heading(doc, "2  Evaluation protocol and safeguards")
    bullet(doc, "Engines, rather than individual rows, are split into 80 development and 20 validation engines.")
    bullet(doc, "Validation evaluates 50%, 70%, and 90% prefixes of held-out trajectories, creating 60 test-like endpoint cases without exposing future cycles as features.")
    bullet(doc, "Features are causal and reset at each engine boundary. Difference-to-initial-reading features use only the first observed point for that engine.")
    bullet(doc, "Model selection ran without opening test_FD001.txt or RUL_FD001.txt. The selected model was then scored once on official endpoints.")
    bullet(doc, "PHM08 is reported alongside RMSE, MAE, R², bias, and error-percentile measures. Its asymmetric penalty is stricter for RUL overprediction.")

    heading(doc, "3  Models implemented")
    add_table(
        doc,
        ["Family", "Role in study", "Result or status"],
        [
            ["Histogram Gradient Boosting", "Selected tabular RUL model", "Best clean validation and final official score"],
            ["CNN, LSTM, CNN-LSTM", "Compact sequence candidates", "CNN-LSTM best neural candidate"],
            ["TinyML-oriented neural regressor", "Small footprint deployment comparison", "Historical official-test comparison"],
            ["Autoencoder", "Early-life reconstruction anomaly signal", "Used for anomaly persistence and triage"],
            ["Hybrid TinyML plus autoencoder", "Joint RUL and anomaly comparison", "Historical deployment comparison"],
            ["Random forest, ridge, mean", "Classical baselines", "Retained for benchmark context"],
        ],
        [4.2, 6.5, 5.8],
    )
    para(doc, "The compact sequence study uses CPU-bounded settings: four threads, batch size 256, hidden width 16, a 20-epoch maximum, and validation early stopping. This kept training suitable for a 16 GB RAM laptop with integrated graphics.")

    heading(doc, "4  Ablation and model comparison")
    para(doc, "Label capping changed the most important behavior: validation PHM08 mean fell from 19.728 for the uncapped 5/20 control to 6.286 with a 125-cycle cap. The selected 10/30 feature variant reduced it again to 2.832.")
    add_table(
        doc,
        ["Candidate", "Validation RMSE", "Validation MAE", "R²", "PHM08 mean"],
        [
            ["HGB uncapped control, 5/20 features", "18.338", "12.409", "0.747", "19.728"],
            ["HGB cap 125, 5/20 features", "16.417", "11.510", "0.797", "6.286"],
            ["HGB cap 125, 10/30 features - selected", "12.773", "8.609", "0.877", "2.832"],
            ["Compact CNN-LSTM", "16.199", "12.488", "0.803", "4.010"],
            ["Compact CNN", "17.525", "13.251", "0.769", "5.428"],
            ["Compact LSTM", "18.592", "13.844", "0.740", "9.125"],
        ],
        [6.8, 3.0, 2.8, 2.0, 2.8],
    )
    if FIGURE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(FIGURE), width=Inches(6.35))
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(6)
        run = caption.add_run("Figure 1. Validation ablation comparison. Lower values are better.")
        run.italic = True
        run.font.size = Pt(8.5)

    heading(doc, "5  Official test result and error profile")
    para(doc, "After selection, the HGB model was retrained using all 100 FD001 training engines. It was evaluated once against the 100 official test endpoints. The near-zero bias and balanced direction counts indicate no material systematic over- or under-estimation in this run.")
    add_table(
        doc,
        ["Metric", "Value", "Metric", "Value"],
        [
            ["RMSE", "13.532", "MAE", "9.799"],
            ["R²", "0.894", "PHM08 sum", "285.222"],
            ["PHM08 mean", "2.852", "Bias, prediction minus actual", "0.157 cycles"],
            ["P90 absolute error", "24.135 cycles", "Maximum absolute error", "38.845 cycles"],
            ["Over / under predictions", "50 / 50", "Official endpoint count", "100"],
        ],
        [4.5, 3.4, 5.0, 3.9],
    )

    heading(doc, "6  On-device anomaly and maintenance workflow")
    para(doc, "An early-life autoencoder learns a reconstruction baseline from development windows. Its 99th-percentile reconstruction-error threshold is used as an unlabeled anomaly signal. It does not have anomaly accuracy metrics because FD001 lacks cycle-level anomaly labels.")
    bullet(doc, "The replay evaluates each incoming record causally and combines RUL, anomaly persistence, data quality, and model version.")
    bullet(doc, "Deterministic triage assigns monitoring, recommendation, planned-maintenance, high-priority, or human-review outcomes.")
    bullet(doc, "Idempotency prevents duplicate work orders for the same engine and severity. The validation replay produced 43 deduplicated mock recommendations.")
    bullet(doc, "An LLM is not on the safety-critical decision path. Any narrative assistant should be limited to explaining already generated evidence and must not override a maintenance decision.")

    heading(doc, "7  Conclusions and recommended next work")
    para(doc, "For FD001, the selected HGB pipeline is the best current choice: it achieved the strongest validation and final official-test metrics with substantially lower training and inference cost than the neural alternatives. The compact neural and autoencoder components remain useful for edge deployment research and anomaly triage, but they should not replace the selected RUL model on the current evidence.")
    bullet(doc, "Validate a deployable prefix-only change-point policy before using it in labeling or in a production decision pipeline.")
    bullet(doc, "Run FD002 and FD004 studies with condition-aware normalization and grouped engine splits.")
    bullet(doc, "Calibrate prediction intervals and maintenance thresholds with domain owners, cost constraints, and false-alarm tolerance.")
    bullet(doc, "Measure TinyML model size, latency, RAM use, quantization error, and update cost on the intended hardware before presenting an on-device deployment claim.")

    heading(doc, "8  Evidence and reproducibility")
    para(doc, "The project contains reproducible runners for the validation-only study, compact sequence study, one-time official-test evaluation, and anomaly-triage replay. The full metrics matrix, model predictions, data audit, metric audit, generated figures, and unit tests are retained in the reports and artifacts directories.")
    para(doc, "Validation completed with 16 passing tests and one skipped optional-dependency scenario. Historical official-test scores are retained as context only; they were produced before this clean selection protocol and were not used to choose the final model.")
    para(doc, "References: NASA C-MAPSS data documentation; Ensarioglu, Inkaya, and Emel (2023), Applied Sciences 13(21), 11893; and the user-provided research-paper materials and implementation notes.")

    properties = doc.core_properties
    properties.title = "NASA C MAPSS FD001 Remaining Useful Life"
    properties.subject = "Final implementation report"
    properties.author = "FD001 RUL project"
    properties.comments = ""
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
